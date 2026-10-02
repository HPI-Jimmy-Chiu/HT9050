from pathlib import Path
from datetime import datetime, timedelta
import re
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

folder = Path(r'D:\00_Weekly Report\2026\20260315')
files = sorted([p for p in folder.glob('RD5_個人週報_*.xlsx') if not p.name.startswith('~$')])

def esc(s):
    if s is None: return ''
    s = str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('|', '&#124;')
    return s.replace('\r\n', '<br>').replace('\n', '<br>').replace('\r', '<br>')

def is_red_color(color):
    if color is None: return False
    color_type = getattr(color, 'type', None)
    if color_type == 'rgb':
        rgb = (getattr(color, 'rgb', '') or '').upper()
        if type(rgb) is str and rgb.endswith('FF0000'): return True
    if color_type == 'indexed':
        return getattr(color, 'indexed', None) in {3, 10}
    return False

def parse_person_date(stem):
    d = datetime(2026, 3, 15)
    for pat in [r'(\d{4})_(\d{2})_(\d{2})', r'(\d{4})-(\d{2})-(\d{2})']:
        m = re.search(pat, stem)
        if m:
            d = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            break
    else:
        m = re.search(r'(\d{8})', stem)
        if m: d = datetime.strptime(m.group(1), '%Y%m%d')
    parts = stem.split('_')
    person = parts[2] if len(parts) >= 3 else stem
    person = re.sub(r'-\d{4}-\d{2}-\d{2}$', '', person)
    person = re.sub(r'\d{8}$', '', person).strip('_-')
    if person == '洪嘉均': person = '洪嘉均'
    return person, d

def normalize_text_dates(text, default_year):
    if text is None:
        return ''

    normalized = str(text)
    normalized = normalized.replace('～', '~').replace('；', '、')

    def repl_datetime(match):
        year, month, day = match.groups()
        return f'{int(year):04d}/{int(month):02d}/{int(day):02d}'

    def repl_compact_range(match):
        start_text, end_text = match.groups()
        start_date = datetime.strptime(start_text, '%Y%m%d')
        end_date = datetime.strptime(end_text, '%Y%m%d')
        return f'{start_date:%Y/%m/%d}~{end_date:%Y/%m/%d}'

    def repl_full_range(match):
        year1, month1, day1, year2, month2, day2 = match.groups()
        return f'{int(year1):04d}/{int(month1):02d}/{int(day1):02d}~{int(year2):04d}/{int(month2):02d}/{int(day2):02d}'

    def repl_shared_year_range(match):
        year, month1, day1, month2, day2 = match.groups()
        return f'{int(year):04d}/{int(month1):02d}/{int(day1):02d}~{int(year):04d}/{int(month2):02d}/{int(day2):02d}'

    def repl_month_day_range(match):
        month1, day1, month2, day2 = match.groups()
        return f'{int(default_year):04d}/{int(month1):02d}/{int(day1):02d}~{int(default_year):04d}/{int(month2):02d}/{int(day2):02d}'

    def repl_month_day_range_with_bullet(match):
        prefix, month1, day1, month2, day2, bullet = match.groups()
        return f'{prefix}{int(default_year):04d}/{int(month1):02d}/{int(day1):02d}~{int(default_year):04d}/{int(month2):02d}/{int(day2):02d}<br>{bullet}'

    def repl_month_day_single(match):
        prefix, month, day = match.groups()
        return f'{prefix}{int(default_year):04d}/{int(month):02d}/{int(day):02d}'

    def repl_month_day_list(match):
        prefix, year, month, day = match.groups()
        return f'{prefix}{int(year):04d}/{int(month):02d}/{int(day):02d}'

    def repl_single(match):
        year, month, day = match.groups()
        return f'{int(year):04d}/{int(month):02d}/{int(day):02d}'

    def repl_compact_single(match):
        value = match.group(1)
        date_value = datetime.strptime(value, '%Y%m%d')
        return f'{date_value:%Y/%m/%d}'

    def repl_numbered_compact_single(match):
        prefix, number, value = match.groups()
        date_value = datetime.strptime(value, '%Y%m%d')
        return f'{prefix}{number}{date_value:%Y/%m/%d}'

    def repl_year_month_days(match):
        parts = re.findall(r'\d{1,4}', match.group(1))
        if len(parts) < 3:
            return match.group(1)
        year = int(parts[0])
        month = int(parts[1])
        days = [f'{year:04d}/{month:02d}/{int(day):02d}' for day in parts[2:]]
        return '、'.join(days)

    def repl_month_day_with_paren(match):
        month, day = match.groups()
        return f'{int(default_year):04d}/{int(month):02d}/{int(day):02d}'

    def repl_seqnum_month_day(match):
        groups = match.groups()
        if len(groups) == 4:
            prefix, seqnum, month, day = groups
            return f'{prefix}{seqnum}. {int(default_year):04d}/{int(month):02d}/{int(day):02d}'
        elif len(groups) == 3:
            seqnum, month, day = groups
            return f'{seqnum}. {int(default_year):04d}/{int(month):02d}/{int(day):02d}'
        return match.group(0)

    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{2})-(\d{1,2})-(\d{1,2})\s+\d{2}:\d{2}:\d{2}(?![A-Za-z0-9_.])', repl_datetime, normalized)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{6})\s*[~-]\s*(20\d{6})(?![A-Za-z0-9_.])', repl_compact_range, normalized)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{2})[./、_-](\d{1,2})[./、_-](\d{1,2})\s*[~-]\s*(20\d{2})[./、_-](\d{1,2})[./、_-](\d{1,2})(?![A-Za-z0-9_.])', repl_full_range, normalized)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{2})/(\d{1,2})/(\d{1,2})\s*[~-]\s*(\d{1,2})/(\d{1,2})(?![A-Za-z0-9_.])', repl_shared_year_range, normalized)
    normalized = re.sub(r'(^|<br>|[\[(、\s>])(\d{1,2})[./](\d{1,2})\s*[~-]\s*(\d{1,2})[./](\d{1,2})([A-Za-z]\.)', repl_month_day_range_with_bullet, normalized, flags=re.MULTILINE)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(\d{1,2})[./](\d{1,2})\s*[~-]\s*(\d{1,2})[./](\d{1,2})(?![A-Za-z0-9_.])', repl_month_day_range, normalized)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{2}[、]\d{1,2}(?:[、]\d{1,2}){2,})(?![A-Za-z0-9_.])', repl_year_month_days, normalized)
    normalized = re.sub(r'(^|<br>|[\[(、\s>])(\d{1,2})[./](\d{1,2})(?=<br>|\)|\]|、|\s|<|$)', repl_month_day_single, normalized, flags=re.MULTILINE)
    normalized = re.sub(r'(^|[\[(、\s])(20\d{2})[./](\d{1,2})[./](\d{1,2})(?=\)|\]|、|\s|$)', repl_month_day_list, normalized)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{2})[./、_-](\d{1,2})[./、_-](\d{1,2})(?![A-Za-z0-9_.])', repl_single, normalized)
    normalized = re.sub(r'(\d{1,2})/(\d{1,2})\([^)]*\)', lambda m: repl_month_day_with_paren(m), normalized)
    normalized = re.sub(r'(^|<br>|>|[\s(])(\d+)\.(\d{1,2})/(\d{1,2})', lambda m: repl_seqnum_month_day(m), normalized, flags=re.MULTILINE)
    normalized = re.sub(r'(\d+)\.(\d{1,2})/(\d{1,2})(?=\s|<|$)', lambda m: repl_seqnum_month_day(m), normalized)
    normalized = re.sub(r'(^|<br>|[\[(、\s])(\d+\.)(20\d{6})(?![A-Za-z0-9_.])', repl_numbered_compact_single, normalized, flags=re.MULTILINE)
    normalized = re.sub(r'(?<![A-Za-z0-9_.])(20\d{6})(?![A-Za-z0-9_.])', repl_compact_single, normalized)

    return normalized

def normalize_cell_display_value(value, default_year):
    if value is None:
        return ''

    if isinstance(value, datetime):
        return value.strftime('%Y/%m/%d')

    year = getattr(value, 'year', None)
    month = getattr(value, 'month', None)
    day = getattr(value, 'day', None)
    if year is not None and month is not None and day is not None:
        return f'{int(year):04d}/{int(month):02d}/{int(day):02d}'

    return normalize_text_dates(str(value), default_year)

def week_text(d):
    mon = d - timedelta(days=d.weekday())
    sun = mon + timedelta(days=6)
    w = ['一', '二', '三', '四', '五', '六', '日']
    return f'{mon:%Y/%m/%d}（週{w[mon.weekday()]}）～ {sun:%Y/%m/%d}（週{w[sun.weekday()]}）'

def cell_to_text(cell):
    value = cell.value
    if isinstance(value, CellRichText):
        return ''.join(item.text if hasattr(item, 'text') else str(item) for item in value)
    return str(value) if value is not None else ''

def get_action_spans(cell):
    value = cell.value
    red_parts = []
    black_parts = []
    has_red = False

    cell_is_red = False
    if hasattr(cell, 'font'):
        cell_is_red = is_red_color(getattr(cell.font, 'color', None))

    if isinstance(value, CellRichText):
        for item in value:
            if isinstance(item, TextBlock):
                txt = item.text or ''
                tb_color = getattr(item.font, 'color', None)
                if tb_color is not None and getattr(tb_color, 'type', None) is not None:
                    red = is_red_color(tb_color)
                else:
                    red = cell_is_red
            else:
                txt = str(item)
                red = cell_is_red
            # if not txt.strip():
            #    continue
            if red:
                has_red = True
                red_parts.append(txt)
            else:
                black_parts.append(txt)
    else:
        txt = '' if value is None else str(value)
        # if txt.strip():
        red = cell_is_red
        if red:
            has_red = True
            red_parts.append(txt)
        else:
            black_parts.append(txt)

    return ("".join(red_parts)).strip(), ("".join(black_parts)).strip(), has_red

def split_and_group_red_text(text, report_date):
    text = normalize_text_dates(text, report_date.year)
    lines = text.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    current_lines = []
    history_lines = []
    
    current_date_block_is_history = False
    mon = report_date - timedelta(days=report_date.weekday())
    
    for line in lines:
        raw_dates = re.findall(r'20\d{2}/\d{2}/\d{2}|20\d{6}', line)
        if raw_dates:
            latest_date = None
            for d_str in raw_dates:
                d_str = d_str.replace('/', '')
                try:
                    d = datetime.strptime(d_str, '%Y%m%d')
                    if not latest_date or d > latest_date:
                        latest_date = d
                except:
                    pass
            if latest_date:
                if latest_date < mon:
                    current_date_block_is_history = True
                else:
                    current_date_block_is_history = False
                    
        if current_date_block_is_history:
            history_lines.append(line)
        else:
            current_lines.append(line)
            
    def group_machines(lines_to_group):
        grouped_result = []
        groups = {}
        other_lines = []
        for l in lines_to_group:
            stripped = l.strip()
            if not stripped: continue
            
            match = re.search(r'(HT-\d+|POLB-\d+|PQ-\d+|HT\d{4}[A-Za-z]*)', l)
            if match and ('改機' in l or '升級' in l):
                machine = match.group(1)
                desc = l.replace(machine, '').strip()
                desc = re.sub(r'^[\d\.\s]+', '', desc).strip()
                if len(desc) > 5:
                    if desc not in groups:
                        groups[desc] = []
                    if machine not in groups[desc]:
                        groups[desc].append(machine)
                    continue
            other_lines.append(l)
                
        if other_lines:
            grouped_result.extend(other_lines)
            
        for desc, machines in groups.items():
            if len(machines) > 1:
                machines_str = '、'.join(sorted(machines))
                grouped_result.append(f'- {machines_str} {desc}')
            else:
                grouped_result.append(f'- {machines[0]} {desc}')
                
        return '\n'.join(grouped_result)

    grouped_current = group_machines(current_lines)
    return grouped_current, '\n'.join(history_lines)

for fp in files:
    try:
        person, report_date = parse_person_date(fp.stem)
        wb = load_workbook(fp, data_only=True, rich_text=True)
        ws = wb.worksheets[0]

        c_dept=1; c_cust=2; c_event=3; c_action=4; c_owner1=5; c_owner2=6; c_plan=7; c_real=8; c_memo=9
        col1_val = str(ws.cell(1, 1).value or "").strip()
        col2_val = str(ws.cell(1, 2).value or "").strip()
        if col1_val == "" and "部門" in col2_val:
            c_dept=2; c_cust=3; c_event=4; c_action=5; c_owner1=6; c_owner2=7; c_plan=8; c_real=9; c_memo=10
        elif "日期" in col1_val and "部門" in col2_val:
            c_dept=2; c_cust=3; c_event=4; c_action=5; c_owner1=6; c_owner2=7; c_plan=8; c_real=9; c_memo=10

        main_rows = []
        history_rows = []
        idx = 1

        for r in range(2, ws.max_row + 1):
            if ws.row_dimensions[r].hidden: continue

            action_cell = ws.cell(r, c_action)
            red_text, black_text, has_red_action = get_action_spans(action_cell)

            if not has_red_action or not red_text:
                continue

            curr_red_text, hist_from_red = split_and_group_red_text(red_text, report_date)
            
            if hist_from_red:
                if black_text:
                    black_text = hist_from_red + "\n" + black_text
                else:
                    black_text = hist_from_red

            if not curr_red_text.strip():
                print(f"[{fp.name}] Row {r} skipped from main: Red text moved to history.")
                continue

            cols = []
            cols.append(str(idx))

            def get_black_span(c_index):
                cell_value = ws.cell(r, c_index).value
                v = normalize_cell_display_value(cell_value, report_date.year)
                if c_index == c_owner1 and not v.strip():
                    v = person
                if not v.strip(): return ""
                return f'<span style="color:#000000">{esc(v)}</span>'

            dept = get_black_span(c_dept)
            if not dept: dept = '<span style="color:#000000">研五</span>'
            cust = get_black_span(c_cust)
            evt = get_black_span(c_event)
            owner1 = get_black_span(c_owner1)
            owner2 = get_black_span(c_owner2)
            plan = get_black_span(c_plan)
            real = get_black_span(c_real)
            memo = get_black_span(c_memo)

            main_act = f'<div id="action-row-{idx}"></div><span style="color:#ff0000">{esc(normalize_text_dates(curr_red_text, report_date.year))}</span>'
            if black_text:
                main_act += f'<br><a href="#action-detail-{idx}">[查看歷史內容]</a>'

            main_rows.append(f"| {idx} | {dept} | {cust} | {evt} | {main_act} | {owner1} | {owner2} | {plan} | {real} | {memo} |")

            if black_text:
                hist_act = f'<span style="color:#000000">{esc(normalize_text_dates(black_text, report_date.year))}</span><br><a href="#action-row-{idx}">[返回主表]</a>'
                hist_row = f'        <tr id="action-detail-{idx}">\n            <td>{idx}</td>\n            <td>{dept}</td>\n            <td>{cust}</td>\n            <td>{evt}</td>\n            <td>{hist_act}</td>\n            <td>{owner1}</td>\n            <td>{owner2}</td>\n            <td>{plan}</td>\n            <td>{real}</td>\n            <td>{memo}</td>\n        </tr>'
                history_rows.append(hist_row)

            idx += 1

        if not main_rows:
            print(f"[{fp.name}] No valid updates found this week.")
            main_rows.append(f"| 1 | <span style=\"color:#000000\">研五</span> |  | 當週無更新 | | <span style=\"color:#000000\">{person}</span> | | | | |")

        lines = [
            "# RD5 個人週報",
            "",
            f"**人員**：{person}  ",
            "**部門**：研五  ",
            f"**週期**：{week_text(report_date)}  ",
            f"**報告日期**：{report_date:%Y-%m-%d}  ",
            "**主要工作區**：D:\\  ",
            "**工作領域**：軟體  ",
            "",
            "## 週報表",
            "",
            "| 序號 | 部門 | 客戶 | 事件/議題 | 行動/解決方案 | 負責人(主) | 負責人(協) | 預定完成日 | 實際完成日 | 備註 |",
            "|------|------|------|-----------|---------------|------------|------------|------------|------------|------|"
        ]
        lines.extend(main_rows)
        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 摘要")
        lines.append("")
        lines.append("### 本週工作總結")
        lines.append("")
        lines.append("系統自動彙整中，待補充。")
        lines.append("")
        lines.append("### 下週建議")
        lines.append("")
        lines.append("無")
        lines.append("")

        if history_rows:
            lines.append("## 行動/解決方案 歷史內容")
            lines.append("")
            lines.append("<table>")
            lines.append("    <colgroup>")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:16%\" />")
            lines.append("        <col style=\"width:40%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:5%\" />")
            lines.append("        <col style=\"width:9%\" />")
            lines.append("    </colgroup>")
            lines.append("    <thead>")
            lines.append("        <tr>")
            lines.append("            <th>序號</th><th>部門</th><th>客戶</th><th>事件/議題</th><th>行動/解決方案</th>")
            lines.append("            <th>負責人(主)</th><th>負責人(協)</th><th>預定完成日</th><th>實際完成日</th><th>備註</th>")
            lines.append("        </tr>")
            lines.append("    </thead>")
            lines.append("    <tbody>")
            lines.extend(history_rows)
            lines.append("    </tbody>")
            lines.append("</table>")
            lines.append("")

        md_path = fp.parent / f'RD5_個人週報_{person}_{report_date:%Y_%m_%d}.md'
        md_path.write_text("\n".join(lines), encoding="utf-8")
        print(f"Generated {md_path.name}")
    except Exception as e:
        print(f"Error processing {fp.name}: {e}")
        import traceback
        traceback.print_exc()
