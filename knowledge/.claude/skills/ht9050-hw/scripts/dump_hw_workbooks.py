# -*- coding: utf-8 -*-
r"""把 HT9050 的三份硬體工作簿全部分頁 dump 成 UTF-8 純文字，方便 diff 與 grep。

用法（一定要用 .venv 的 python，PATH 上的 python 是 Store 殼，會靜默 exit 49）：

    & d:\HT9045\.venv\Scripts\python.exe dump_hw_workbooks.py [--out DIR] [--src DIR]

輸出：
    <out>/HP-9050開發機資料.txt        (.xlsx, openpyxl)
    <out>/HP-9050元件代碼-軟體.txt     (.xls,  xlrd)
    <out>/HP-9050元件代碼-電控.txt     (.xls,  xlrd)
    <out>/_diff-軟體-vs-電控.txt       兩份 .xls 的逐格差異

硬體送新版工作簿來的時候：先跑這支、再對舊的 dump 做 diff，不要用眼睛比。
"""
import argparse
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_SRC = os.path.join(os.path.dirname(HERE), "docs")   # skill 的 docs/

XLSX = "HP-9050開發機資料-20260717.xlsx"
XLS_SW = "HP-9050機構類元件代碼-20260922-2-軟體.xls"
XLS_EC = "HP-9050機構類元件代碼-20260923-1-電控.xls"


def find(src, prefix, suffix):
    """檔名帶日期，所以用前後綴找，換版本不用改腳本。"""
    hits = sorted(f for f in os.listdir(src)
                  if f.startswith(prefix) and f.endswith(suffix) and not f.startswith("~$"))
    if not hits:
        return None
    return os.path.join(src, hits[-1])       # 取字典序最後一個 == 日期最新


def dump_xlsx(path, out_path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    with io.open(out_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("######## FILE: %s\n\n" % os.path.basename(path))
        for ws in wb.worksheets:
            fh.write("=== SHEET: %s (%d x %d)\n" % (ws.title, ws.max_row, ws.max_column))
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    fh.write(" | ".join(cells) + "\n")
            fh.write("\n")


def read_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    sheets = []
    for ws in wb.sheets():
        rows = [[str(ws.cell_value(i, j)).strip() for j in range(ws.ncols)]
                for i in range(ws.nrows)]
        sheets.append((ws.name, rows))
    return sheets


def dump_xls(path, out_path):
    sheets = read_xls(path)
    with io.open(out_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("######## FILE: %s\n\n" % os.path.basename(path))
        for name, rows in sheets:
            fh.write("=== SHEET: %s (%d x %d)\n" % (name, len(rows), len(rows[0]) if rows else 0))
            for cells in rows:
                if any(cells):
                    fh.write(" | ".join(cells) + "\n")
            fh.write("\n")
    return sheets


def diff_xls(sw, ec, out_path):
    """逐格比對兩份 .xls 的第一個分頁；回傳不同的列數。"""
    a = sw[0][1] if sw else []
    b = ec[0][1] if ec else []
    n = 0
    with io.open(out_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("軟體版 vs 電控版，分頁 %s\n\n" % (sw[0][0] if sw else "?"))
        for i in range(max(len(a), len(b))):
            ra = a[i] if i < len(a) else []
            rb = b[i] if i < len(b) else []
            if ra == rb:
                continue
            n += 1
            for j in range(max(len(ra), len(rb))):
                x = ra[j] if j < len(ra) else ""
                y = rb[j] if j < len(rb) else ""
                if x != y:
                    fh.write("row%-4d col%-3d 軟體[%s] 電控[%s]\n" % (i + 1, j + 1, x, y))
        fh.write("\n不同的列數: %d\n" % n)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=DEFAULT_SRC, help="工作簿所在目錄（預設：skill 的 docs/）")
    ap.add_argument("--out", default=None, help="輸出目錄（預設：--src 底下的 _dump）")
    args = ap.parse_args()

    out = args.out or os.path.join(args.src, "_dump")
    if not os.path.isdir(out):
        os.makedirs(out)

    p_xlsx = find(args.src, "HP-9050開發機資料", ".xlsx")
    p_sw = find(args.src, "HP-9050機構類元件代碼", "軟體.xls")
    p_ec = find(args.src, "HP-9050機構類元件代碼", "電控.xls")

    rc = 0
    if p_xlsx:
        dump_xlsx(p_xlsx, os.path.join(out, "HP-9050開發機資料.txt"))
        print("OK  %s" % os.path.basename(p_xlsx))
    else:
        print("找不到 開發機資料 .xlsx", file=sys.stderr); rc = 1

    sw = ec = None
    if p_sw:
        sw = dump_xls(p_sw, os.path.join(out, "HP-9050元件代碼-軟體.txt"))
        print("OK  %s" % os.path.basename(p_sw))
    else:
        print("找不到 元件代碼-軟體 .xls", file=sys.stderr); rc = 1

    if p_ec:
        ec = dump_xls(p_ec, os.path.join(out, "HP-9050元件代碼-電控.txt"))
        print("OK  %s" % os.path.basename(p_ec))
    else:
        print("找不到 元件代碼-電控 .xls", file=sys.stderr); rc = 1

    if sw and ec:
        n = diff_xls(sw, ec, os.path.join(out, "_diff-軟體-vs-電控.txt"))
        print("軟體 vs 電控：%d 列不同（20260922/20260923 這組基準是 29）" % n)

    print("輸出目錄: %s" % out)
    return rc


if __name__ == "__main__":
    sys.exit(main())
