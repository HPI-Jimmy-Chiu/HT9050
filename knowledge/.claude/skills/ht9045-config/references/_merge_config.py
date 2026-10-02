# -*- coding: utf-8 -*-
"""
Merge ConfigList_20260326_KevinCheng.xlsx  +
      HT-9046 版本899-Configuration 全功能說明_ 修改.xlsx
→ config-full-list.md
"""
import re
from openpyxl import load_workbook

PATH1 = r'D:\HT9045\.github\skills\ht9045-config\references\ConfigList_20260326_KevinCheng.xlsx'
PATH2 = r'D:\HT9045\.github\skills\ht9045-config\references\HT-9046 版本899-Configuration 全功能說明_ 修改.xlsx'
OUT   = r'D:\HT9045\.github\skills\ht9045-config\references\config-full-list.md'

# Group letter → display name
GROUP_NAMES = {
    'A': 'A — Function（自動化 / 流程控制）',
    'B': 'B — Report（報表 / 紀錄）',
    'C': 'C — Hardware（硬體選配）',
    'D': 'D — Index（索引手臂）',
    'E': 'E — In/Out Arm（進出手臂）',
    'F': 'F — Shuttle（梭式機構）',
    'G': 'G — Visible（UI 顯示）',
    'I': 'I — Tester（測試介面）',
    'L': 'L — Temperature（溫度控制）',
    'M': 'M — Monitor（Monitor 強制）',
    'N': 'N — Network（網路 / 上傳）',
    'O': 'O — Count（計數 / 統計）',
    'P': 'P — Tray（盤子系統）',
}
GROUP_ORDER = list('ABCDEFGILMNOP')

_CODE_RE = re.compile(r'\[([A-Z]\d{2}(?:[_\-]\d+)*)\]', re.IGNORECASE)

def extract_code(text):
    """從字串中提取第一個功能代碼，如 A01、A01-1、N14-5 等"""
    if not text:
        return None
    m = _CODE_RE.search(str(text))
    if not m:
        return None
    return m.group(1).upper().replace('_', '-')

def clean(v):
    if v is None:
        return ''
    s = str(v).strip()
    # 去掉 Excel 因欄位過窄顯示的 ####
    if s.startswith('#'):
        return ''
    return s

# ── 讀 File 1 ────────────────────────────────────────────────
# Columns: Name(1) ECID(2) Type(3) Function(4) Sub-Function(5) Description(6) ChineseDesc(7)
print('Reading File1...')
wb1 = load_workbook(PATH1, read_only=True, data_only=True)
ws1 = wb1['工作表1']

# items[code] = {ecid, type_, en_desc, zh_desc, group_letter}
items = {}   # key = code like 'A01', 'A01-1'
code_order = []   # maintain order

for ri, row in enumerate(ws1.iter_rows(min_row=2, max_row=ws1.max_row, values_only=True)):
    fn_col  = clean(row[3])   # Function
    sub_col = clean(row[4])   # Sub-Function
    ecid    = clean(row[1])
    type_   = clean(row[2])
    en_desc = clean(row[5])
    zh_desc = clean(row[6])
    name    = clean(row[0])

    # Determine group letter from Name col (e.g. "A [ Function ]")
    group_letter = ''
    if name:
        m = re.match(r'^([A-Z])\s*\[', name)
        if m:
            group_letter = m.group(1)

    # Try to get code from Function first, then Sub-Function
    code = extract_code(fn_col) or extract_code(sub_col)
    if not code:
        continue

    # Always use code's first letter as ground truth (Name column may have errors)
    group_letter = code[0].upper()

    # English description: prefer Function col content (strip code), else Sub-Function
    def strip_code_prefix(s):
        return _CODE_RE.sub('', s).strip()

    en = strip_code_prefix(fn_col) if fn_col else strip_code_prefix(sub_col)
    if not en and sub_col:
        en = strip_code_prefix(sub_col)

    if code not in items:
        items[code] = {
            'group': group_letter,
            'ecid': ecid,
            'type': type_,
            'en': en,
            'zh': zh_desc,
            'note': '',
        }
        code_order.append(code)
    else:
        # merge: fill blanks
        d = items[code]
        if not d['ecid'] and ecid:   d['ecid'] = ecid
        if not d['type'] and type_:  d['type']  = type_
        if not d['en']   and en:     d['en']    = en
        if not d['zh']   and zh_desc: d['zh']   = zh_desc
        if not d['group'] and group_letter: d['group'] = group_letter

wb1.close()
print(f'  File1: {len(items)} unique codes')

# ── 讀 File 2 ────────────────────────────────────────────────
# Columns: 項目(1) 功能分類(2) 功能說明(3) 備註(4)
print('Reading File2...')
wb2 = load_workbook(PATH2, read_only=True, data_only=True)
ws2 = wb2['HT9045HW config-PM']

cur_group = ''
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, values_only=True):
    item_col = clean(row[0])
    fn_col2  = clean(row[1])
    zh2      = clean(row[2])
    note2    = clean(row[3])

    if item_col and item_col != '項目':
        # Map English group names to letters
        _gmap = {
            'function': 'A', 'report': 'B', 'hardware': 'C',
            'index': 'D', 'in/out arm': 'E', 'shuttle': 'F',
            'visible': 'G', 'tester': 'I', 'temperature': 'L',
            'monitor': 'M', 'network': 'N', 'count': 'O', 'tray': 'P',
        }
        cur_group = _gmap.get(item_col.lower(), cur_group)

    code2 = extract_code(fn_col2)
    if not code2:
        continue

    # English name: strip code from fn_col2
    en2 = _CODE_RE.sub('', fn_col2).strip()
    # Remove leading spaces
    en2 = re.sub(r'\s{2,}', ' ', en2).strip()

    if code2 in items:
        d = items[code2]
        if not d['zh'] and zh2:    d['zh']   = zh2
        if not d['note'] and note2 and note2 != '*': d['note'] = note2
        if not d['en'] and en2:    d['en']   = en2
        # Always override group with code's first letter
        d['group'] = code2[0].upper()
    else:
        # Code exists in File2 but not File1
        items[code2] = {
            'group': code2[0].upper(),  # always from code itself
            'ecid': '',
            'type': '',
            'en': en2,
            'zh': zh2,
            'note': note2 if note2 != '*' else '',
        }
        code_order.append(code2)

wb2.close()
print(f'  After File2 merge: {len(items)} unique codes')

# ── Sort codes ───────────────────────────────────────────────
# Natural sort: group letter first, then numeric parts
def code_sort_key(code):
    parts = re.split(r'[-_]', code)
    nums = []
    for p in parts[1:]:   # skip letter prefix
        try:
            nums.append(int(p))
        except:
            nums.append(0)
    letter = parts[0][0].upper() if parts else 'Z'
    main_num = int(parts[0][1:]) if len(parts[0]) > 1 and parts[0][1:].isdigit() else 0
    return (GROUP_ORDER.index(letter) if letter in GROUP_ORDER else 99, main_num, nums)

sorted_codes = sorted(items.keys(), key=code_sort_key)

# ── Build Markdown ───────────────────────────────────────────
print('Writing Markdown...')
lines = []
lines.append('# HT9045 Config 全功能清單')
lines.append('')
lines.append('> 整合來源：')
lines.append('> - `ConfigList_20260326_KevinCheng.xlsx`（ECID / Type / 英文說明）')
lines.append('> - `HT-9046 版本899-Configuration 全功能說明_ 修改.xlsx`（中文說明 / 備註）')
lines.append('')
lines.append('---')
lines.append('')

# TOC
lines.append('## 目錄')
lines.append('')
for letter in GROUP_ORDER:
    gname = GROUP_NAMES.get(letter, letter)
    anchor = gname.lower().replace(' ', '-').replace('（', '').replace('）', '').replace('/', '').replace('—', '').strip('-')
    lines.append(f'- [{gname}](#{anchor})')
lines.append('')
lines.append('---')
lines.append('')

cur_group = ''
for code in sorted_codes:
    d = items[code]
    g = d['group']
    if not g or g not in GROUP_ORDER:
        continue

    # Group header
    if g != cur_group:
        cur_group = g
        gname = GROUP_NAMES.get(g, g)
        lines.append(f'## {gname}')
        lines.append('')
        lines.append('| 代碼 | ECID | Type | 功能說明（英文）| 功能說明（中文）| 備註 |')
        lines.append('|------|------|------|----------------|----------------|------|')

    def md_cell(s):
        return s.replace('|', '｜').replace('\n', ' ').strip()

    ecid  = md_cell(d['ecid'])
    type_ = md_cell(d['type'])
    en    = md_cell(d['en'])
    zh    = md_cell(d['zh'])
    note  = md_cell(d['note'])

    # Determine indentation for sub-functions
    depth = code.count('-')
    indent = '&nbsp;&nbsp;' * depth
    code_disp = f'{indent}[{code}]'

    lines.append(f'| {code_disp} | {ecid} | {type_} | {en} | {zh} | {note} |')

lines.append('')
lines.append('---')
lines.append(f'')
lines.append(f'*自動產生於 2026-04-01，來源：兩份 Excel 整合*')

with open(OUT, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))

print(f'Done → {OUT}')
